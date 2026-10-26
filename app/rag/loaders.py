"""Multi-format document loaders for the knowledge base (json, csv, txt/md, docx, pdf).
Folder name = category (places, food, stay, transport, events, policy, faq, guides, safety, activities).
Pure Python, no network: easy to unit-test. Pinecone is only touched in ingest.py."""
# import csv, hashlib, json, pathlib, re
# from typing import Iterator

# MAX_CHARS = 900  # target chunk size for prose


# def _title(path: pathlib.Path) -> str:
#     return path.stem.replace("_", " ").title()


# def _flatten(d: dict) -> str:
#     return "; ".join(f"{k.replace('_', ' ')}: {v}" for k, v in d.items() if v not in ("", None))


# def _split_prose(text: str, max_chars: int = MAX_CHARS) -> list[str]:
#     """Paragraph-aware chunking: merge short paragraphs, split long ones on sentence ends."""
#     paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
#     out, buf = [], ""
#     for p in paras:
#         if len(p) > max_chars:
#             if buf:
#                 out.append(buf); buf = ""
#             sent, cur = re.split(r"(?<=[.!?])\s+", p), ""
#             for s in sent:
#                 if len(cur) + len(s) > max_chars and cur:
#                     out.append(cur.strip()); cur = ""
#                 cur += s + " "
#             if cur.strip():
#                 out.append(cur.strip())
#         elif len(buf) + len(p) + 2 > max_chars:
#             out.append(buf); buf = p
#         else:
#             buf = f"{buf}\n\n{p}".strip()
#     if buf:
#         out.append(buf)
#     return out


# def load_json(path):
#     data = json.loads(path.read_text(encoding="utf-8"))
#     items = data if isinstance(data, list) else data.get("items", [data])
#     for it in items:
#         yield (it["text"] if isinstance(it, dict) and "text" in it else _flatten(it)) if isinstance(it, dict) else str(it)


# def load_csv(path):
#     with path.open(encoding="utf-8", newline="") as f:
#         for row in csv.DictReader(f):
#             yield _flatten(row)


# def load_txt(path):
#     yield from _split_prose(path.read_text(encoding="utf-8"))


# def _style_id(p) -> str:
#     """Style id read from XML (works for Word-made and docx-js-made files alike)."""
#     ppr = p._p.pPr
#     return (ppr.pStyle.val if ppr is not None and ppr.pStyle is not None else "").lower()


# def load_docx(path):
#     from docx import Document  # python-docx
#     doc = Document(str(path))
#     heading, buf = "", []
#     sections = []
#     for p in doc.paragraphs:
#         if not p.text.strip():
#             continue
#         if _style_id(p).startswith(("heading", "title")):
#             if buf:
#                 sections.append((heading, "\n".join(buf)))
#             heading, buf = p.text.strip(), []
#         else:
#             buf.append(p.text.strip())
#     if buf:
#         sections.append((heading, "\n".join(buf)))
#     for t in doc.tables:                                   # tables become 'col: value' rows
#         hdr = [c.text.strip() for c in t.rows[0].cells]
#         for r in t.rows[1:]:
#             sections.append(("Table", _flatten(dict(zip(hdr, [c.text.strip() for c in r.cells])))))
#     for h, body in sections:
#         for chunk in _split_prose(body):
#             yield f"{h}: {chunk}" if h else chunk


# def load_pdf(path):
#     from pypdf import PdfReader
#     reader = PdfReader(str(path))
#     for page_no, page in enumerate(reader.pages, 1):
#         text = (page.extract_text() or "").strip()
#         text = re.sub(r"Goa Travel Assistant sample knowledge base[^\n]*\s*Page \d+", "", text).strip()  # drop footer
#         for chunk in _split_prose(text):
#             yield f"{chunk} (page {page_no})"


# LOADERS = {".json": load_json, ".csv": load_csv, ".txt": load_txt, ".md": load_txt,
#            ".docx": load_docx, ".pdf": load_pdf}


# def iter_chunks(root: str = "data/knowledge") -> Iterator[dict]:
#     root_p = pathlib.Path(root)
#     for path in sorted(root_p.rglob("*")):
#         loader = LOADERS.get(path.suffix.lower())
#         if not loader or not path.is_file():
#             continue
#         category = path.parent.name if path.parent != root_p else "general"
#         for i, text in enumerate(loader(path)):
#             text = re.sub(r"\s+", " ", text).strip()
#             if len(text) < 20:
#                 continue
#             uid = hashlib.sha1(f"{path.name}:{i}".encode()).hexdigest()[:16]
#             yield {"id": f"{category}-{uid}", "text": f"[{_title(path)}] {text}",
#                    "metadata": {"category": category, "source": path.name, "format": path.suffix.lstrip("."),
#                                 "chunk": i}}


import csv,hashlib,json,pathlib
from docx import Document
from pypdf import PdfReader
from openpyxl import load_workbook
def read_chunks(path:pathlib.Path) -> list[str]:
    
    extension = path.suffix.lower()

    if extension == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        items = data if isinstance(data,list) else [data]
        return [items.get("text") or json.dumps(item) for item in items]

    if extension == ".csv":
        with path.open(encoding="utf-8",newline="") as f:
            return [ "; ".join(f"{k}: {v}" for k,v in row.items() if v)  for row in csv.DictReader(f)]

    if extension in (".txt",".md"):
        text = path.read_text(encoding="utf-8")
        return [ p.strip() for p in text.split("\n\n") if p.strip()]

    if extension == ".docx":
        heading , out = "" , []

        for p in Document(str(path)).paragraphs:
            text = p.text.strip()
            if not text:
                continue

            if p.style.name.startswith("Heading"):
                heading = text
            else:
                out.append(f"{heading}: {text}" if heading else text)

        return out

    if extension == ".pdf":
        pages = PdfReader(str(path)).pages
        return [f"{(pages.extract_test() or '').strip()} (page {n})"
                for n,page in enumerate(pages,1)
                if (page.extract_text() or '').strip()
                ]

    if extension == ".xlxs":
        wb = load_workbook(str(path),read_only=True,data_only=True)
        out = []
        for sheet in wb.worksheets:
            rows = sheet.iter_rows(values_only=True)
            header = next(rows,None)
            if not header:
                continue
            names = [str(h) for h in header]
            for row in rows:
                pairs = [f"{k} : {v}" for k,v in zip(names,row) if v not in (None,"")]
                if pairs:
                    out.append(f"{sheet.title} | " + "; ".join(pairs))

        return out

    return []

def iter_chunks(root: str="data/knowledge"):
    root = pathlib.Path(root)

    for path in sorted(root.rglob("*")):
        if not path.is_file():
            continue

        category = path.parent.name if path.parent != root else "general"
        title = path.stem.replace("_"," ").title()

        for i , text in enumerate(read_chunks(path)):
            text = " ".join(text.split())
            if len(text) < 30:
                continue

            uid = hashlib.sha1(f"{path.name}: {i}".encode()).hexdigest()[:16]
            yield {
                "id": f"{category}--{uid}",
                "text":f"[{title}] {text}",
                "metadata":{
                    "category":category,
                    "source":path.name,
                    "format":path.suffix.lstrip("."),
                    "chunk":i
                }
            }